"""A follow-up to the team review pack: more detail for every row the team left Unsure.

Usage:
    python -m atdr.scripts.team_review_followup build
    python -m atdr.scripts.team_review_followup read --pack "returned follow-up.xlsx"

build  writes .tmp/team_review/ATDR_team_review_followup.xlsx (anonymized like the pack) and the
       private followup_key.json. It gives the detail the reviewer asked for: each firewall threat
       record with its signature, the biggest uploads and the individual flows for the model rows;
       the source's flows around each blind-check log; and per-device behaviour for the
       unidentified-app labels. Every Unsure row is included, so the model review stays blind.
read   merges the answers into the pack's outputs under .tmp/team_review/ (keeping copies of the
       earlier files) and writes label_decisions_followup.json. It changes nothing in ATDR.

Needs openpyxl, like the pack.
"""

from __future__ import annotations

import argparse
import csv
import json
import shutil
import sqlite3
import statistics
from collections import Counter
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any

from atdr.app.db.database import SessionLocal
from atdr.app.services.blind_check_service import score_blind_check
from atdr.app.services.team_review_service import (
    blind_comparison,
    human_checked_decisions,
    merge_followup,
    pattern_condition_failures,
    signoff_problems,
)
from atdr.scripts.anonymized_review_files import (
    BLIND,
    BLIND_COLUMNS,
    KEY,
    LIVE_DB,
    MODEL_COLUMNS,
    PRIVACY_LINES,
    SIGNOFF,
    _anonymize_rows,
    _write_workbook,
)
from atdr.scripts.team_review_pack import CONFIDENCE, DECISIONS, MODEL_ROUND, PACK_DIR, PACK_KEY, _names, _sheet_rows, _write_csv

FOLLOWUP = PACK_DIR / "ATDR_team_review_followup.xlsx"
FOLLOWUP_KEY = PACK_DIR / "followup_key.json"
HOLDOUT = BLIND / "holdout.db"
SHEET_A, SHEET_B, SHEET_D = "A2. Unsure model rows", "B2. Unsure blind logs", "D2. Unidentified-app devices"
DEVICE_CHOICES = ["Normal but unusual (not an attack)", "Keep as a threat", "Unsure"]
NOT_IN_LOGS = "Process names and TLS certificate details are not recorded in firewall logs."
FLOW_LINES = 25

MODEL_DETAIL_COLUMNS = [column for column in MODEL_COLUMNS if column[3] == "log"] + [
    ("earlier_note", "Earlier note (why it was Unsure)", 40, "log"),
    ("threats", "Firewall threat records (time, signature, severity, action)", 50, "log"),
    ("uploads", "Biggest uploads (flow; bytes sent / received; seconds)", 50, "log"),
    ("destinations", "Traffic per destination (flows; sent / received; country)", 50, "log"),
    ("flows", f"First {FLOW_LINES} flows", 70, "log"),
    ("decision", "YOUR DECISION", 18, "input"), ("confidence", "Confidence", 12, "input"), ("note", "Note (why)", 40, "input"),
]
BLIND_DETAIL_COLUMNS = [column for column in BLIND_COLUMNS if column[3] == "log"] + [
    ("earlier_note", "Earlier note (why it was Unsure)", 40, "log"),
    ("threats", "Same source's firewall threat records, 5 min around", 44, "log"),
    ("flows", f"Same source's flows, 2.5 min either side (first {FLOW_LINES})", 70, "log"),
    ("decision", "YOUR DECISION", 18, "input"), ("confidence", "Confidence", 12, "input"), ("note", "Note (why)", 40, "input"),
]
DEVICE_COLUMNS = [
    ("device", "Device", 16, "log"), ("labels", "Labels in question (current)", 30, "log"), ("span", "Seen (20 May)", 16, "log"),
    ("flows", "Flows", 8, "log"), ("udp_share", "UDP share", 9, "log"), ("unidentified_share", "Unidentified-app share", 11, "log"),
    ("destinations", "Different destinations", 11, "log"), ("ports", "Different destination ports", 11, "log"),
    ("bytes", "Bytes sent / received", 20, "log"), ("no_reply", "Flows with nothing received", 11, "log"),
    ("top_ports", "Top destination ports (flows)", 30, "log"), ("top_destinations", "Top destinations (flows, country)", 44, "log"),
    ("cadence", "Regularity to the top destination", 30, "log"), ("endings", "How sessions ended", 26, "log"),
    ("threats", "Firewall threat records from the device", 22, "log"), ("alerts", "ATDR alerts on it", 24, "log"),
    ("decision", "YOUR DECISION FOR THIS DEVICE'S LABELS", 26, "input"), ("note", "Note (why)", 40, "input"),
]

FLOW_QUERY = """
    SELECT generated_time, log_type, subtype, dst_ip, dst_port, protocol, app, action, bytes_sent, bytes_received,
           elapsed_time, session_end_reason, dst_country, parsed_json
    FROM normalized_logs WHERE src_ip = ? AND generated_time >= ? AND generated_time < ? ORDER BY generated_time, id
"""


def _connect(path: Path) -> sqlite3.Connection:
    return sqlite3.connect(f"file:{path.as_posix()}?mode=ro", uri=True)


def _flow_rows(connection: sqlite3.Connection, source: str, start: datetime, end: datetime) -> list[dict[str, Any]]:
    rows = []
    for when, log_type, subtype, dst, port, proto, app, action, sent, received, elapsed, ended, country, parsed in connection.execute(
        FLOW_QUERY, (source, start.strftime("%Y-%m-%d %H:%M:%S"), end.strftime("%Y-%m-%d %H:%M:%S"))
    ):
        details = json.loads(parsed) if parsed else {}
        rows.append({"time": str(when)[11:19], "log_type": (log_type or "").upper(), "subtype": subtype or "", "dst": dst or "-",
                     "port": port, "proto": proto or "-", "app": app or "-", "action": action or "-", "sent": sent or 0,
                     "received": received or 0, "elapsed": elapsed, "ended": ended or "-", "country": country or "-",
                     "threat": details.get("parsed_threat_name") or "", "severity": details.get("parsed_threat_severity") or ""})
    return rows


def _flow_line(flow: dict[str, Any]) -> str:
    return (f"{flow['time']} -> {flow['dst']}:{flow['port']} {flow['proto']} {flow['app']} {flow['action']}, "
            f"sent {flow['sent']:,} / received {flow['received']:,} bytes, {flow['elapsed'] if flow['elapsed'] is not None else '-'} s, {flow['ended']}")


def _threat_lines(flows: list[dict[str, Any]]) -> str:
    threats = [flow for flow in flows if flow["log_type"] == "THREAT"]
    if not threats:
        return "None"
    return "\n".join(f"{flow['time']} {flow['threat'] or '(no name)'} ({flow['severity'] or '-'}, {flow['subtype'] or '-'}) "
                     f"{flow['action']} -> {flow['dst']}:{flow['port']}" for flow in threats[:15])


def _traffic(flows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [flow for flow in flows if flow["log_type"] != "THREAT"]


def _uploads(flows: list[dict[str, Any]]) -> str:
    biggest = sorted(_traffic(flows), key=lambda flow: flow["sent"], reverse=True)[:5]
    return "\n".join(_flow_line(flow) for flow in biggest) or "None"


def _destinations(flows: list[dict[str, Any]]) -> str:
    totals: dict[str, list[int]] = {}
    country: dict[str, str] = {}
    for flow in _traffic(flows):
        entry = totals.setdefault(flow["dst"], [0, 0, 0])
        entry[0] += 1
        entry[1] += flow["sent"]
        entry[2] += flow["received"]
        country[flow["dst"]] = flow["country"]
    ranked = sorted(totals.items(), key=lambda item: (item[1][1], item[1][0]), reverse=True)[:8]
    return "\n".join(f"{dst}: {count} flows; {sent:,} / {received:,} bytes; {country[dst]}" for dst, (count, sent, received) in ranked) or "None"


def _flows_text(flows: list[dict[str, Any]]) -> str:
    traffic = _traffic(flows)
    lines = [_flow_line(flow) for flow in traffic[:FLOW_LINES]]
    if len(traffic) > FLOW_LINES:
        lines.append(f"... and {len(traffic) - FLOW_LINES} more flows")
    return "\n".join(lines) or "None"


def _model_rows(names: Any) -> tuple[list[dict[str, Any]], list[str]]:
    decisions = {row["review_id"]: row for row in csv.DictReader((PACK_DIR / "model_review_decisions.csv").open(encoding="utf-8"))}
    unsure = sorted(review for review, row in decisions.items() if row["decision"] == "Unsure")
    key = {entry["review_id"]: entry for entry in json.loads((MODEL_ROUND / "review_key.json").read_text(encoding="utf-8"))["reviews"]}
    samples = {row["review_id"]: row for row in json.loads((MODEL_ROUND / "review_sample.json").read_text(encoding="utf-8"))}
    rows = []
    with _connect(HOLDOUT) as connection:
        for review in unsure:
            start = datetime.fromisoformat(key[review]["window"])
            flows = _flow_rows(connection, key[review]["source"], start, start + timedelta(minutes=5))
            rows.append({**samples[review], "earlier_note": decisions[review].get("note", ""), "threats": _threat_lines(flows),
                         "uploads": _uploads(flows), "destinations": _destinations(flows), "flows": _flows_text(flows)})
    return _anonymize_rows(rows, names), unsure


def _blind_rows(names: Any) -> tuple[list[dict[str, Any]], list[str]]:
    decisions = {row["sample_id"]: row for row in csv.DictReader((PACK_DIR / "blind_human_decisions.csv").open(encoding="utf-8"))}
    unsure = sorted(sample for sample, row in decisions.items() if row["decision"] == "Unsure")
    samples = {row["sample_id"]: row for row in json.loads((BLIND / "sample.json").read_text(encoding="utf-8"))}
    log_of = {entry["sample_id"]: entry["log_id"] for entry in json.loads((BLIND / "key.json").read_text(encoding="utf-8"))["samples"]}
    rows = []
    with _connect(HOLDOUT) as connection:
        for sample in unsure:
            source, when = connection.execute("SELECT src_ip, generated_time FROM normalized_logs WHERE id = ?", (log_of[sample],)).fetchone()
            moment = datetime.fromisoformat(str(when)[:19])
            around = _flow_rows(connection, source, moment - timedelta(seconds=150), moment + timedelta(seconds=150))
            threats = _flow_rows(connection, source, moment - timedelta(minutes=5), moment + timedelta(minutes=5))
            rows.append({**samples[sample], "earlier_note": decisions[sample].get("note", ""), "threats": _threat_lines(threats),
                         "flows": _flows_text(around)})
    return _anonymize_rows(rows, names), unsure


def _device_rows(names: Any) -> tuple[list[dict[str, Any]], dict[str, list[int]]]:
    key = json.loads(PACK_KEY.read_text(encoding="utf-8"))
    group = next(pattern for pattern in key["patterns"].values() if pattern["key"] == "unidentified_app")
    by_device: dict[str, list[int]] = {}
    rows = []
    with _connect(LIVE_DB) as connection:
        placeholders = ",".join("?" for _ in group["log_ids"])
        labels: dict[str, Counter] = {}
        for log_id, source in connection.execute(f"SELECT id, src_ip FROM normalized_logs WHERE id IN ({placeholders})", group["log_ids"]):
            by_device.setdefault(source, []).append(int(log_id))
        for source, log_ids in by_device.items():
            marks = ",".join("?" for _ in log_ids)
            labels[source] = Counter(
                f"{label} / {attack}" for label, attack in connection.execute(
                    f"""SELECT l.label, l.attack_type FROM ml_labels l WHERE l.reviewed = 1 AND l.log_id IN ({marks})
                        AND l.id = (SELECT MAX(id) FROM ml_labels WHERE log_id = l.log_id AND reviewed = 1)""", log_ids))
            flows = _flow_rows(connection, source, datetime(2000, 1, 1), datetime(2100, 1, 1))
            traffic = _traffic(flows)
            total = len(traffic) or 1
            ports = Counter(f"{flow['proto']}/{flow['port']}" for flow in traffic)
            destinations = Counter(flow["dst"] for flow in traffic)
            country = {flow["dst"]: flow["country"] for flow in traffic}
            top = destinations.most_common(1)[0][0] if destinations else None
            times = sorted(datetime.strptime(flow["time"], "%H:%M:%S") for flow in traffic if flow["dst"] == top)
            gaps = [(later - earlier).total_seconds() for earlier, later in zip(times, times[1:]) if later > earlier]
            alerts = connection.execute(
                f"""SELECT DISTINCT a.id, a.severity, a.alert_type FROM alerts a JOIN alert_evidence e ON e.alert_id = a.id
                    WHERE e.normalized_log_id IN ({marks})""", log_ids).fetchall()
            rows.append({
                "device": source,
                "labels": ", ".join(f"{text} ({count})" for text, count in labels[source].most_common()),
                "span": f"{traffic[0]['time']}-{traffic[-1]['time']}" if traffic else "-",
                "flows": len(traffic),
                "udp_share": f"{sum(flow['proto'] == 'udp' for flow in traffic) / total:.0%}",
                "unidentified_share": f"{sum(flow['app'].startswith('unknown') or flow['app'] in ('incomplete', 'insufficient-data') for flow in traffic) / total:.0%}",
                "destinations": len(destinations),
                "ports": len(ports),
                "bytes": f"{sum(flow['sent'] for flow in traffic):,} / {sum(flow['received'] for flow in traffic):,}",
                "no_reply": f"{sum(flow['received'] == 0 for flow in traffic) / total:.0%}",
                "top_ports": ", ".join(f"{port} ({count})" for port, count in ports.most_common(5)),
                "top_destinations": "\n".join(f"{dst} ({count}, {country[dst]})" for dst, count in destinations.most_common(5)),
                "cadence": (f"{len(times)} connections, median gap {statistics.median(gaps):.1f} s" if len(gaps) >= 2 else "too few connections"),
                "endings": ", ".join(f"{ended} ({count})" for ended, count in Counter(flow["ended"] for flow in traffic).most_common(3)),
                "threats": _threat_lines(flows),
                "alerts": ", ".join(f"#{alert_id} {severity} ({alert_type.replace('_', ' ')})" for alert_id, severity, alert_type in alerts) or "None",
            })
    return _anonymize_rows(rows, names), by_device


def build() -> None:
    names = _names()
    model_rows, model_ids = _model_rows(names)
    blind_rows, blind_ids = _blind_rows(names)
    devices, by_device = _device_rows(names)
    device_key = {row["device"]: by_device[real] for row, real in zip(devices, by_device)}
    guide = [
        ("ATDR team review follow-up: the rows left Unsure", True), ("", False),
        ("The review left some rows Unsure and asked for more detail. This file has that detail for every Unsure row. It is", False),
        ("the same data as before, from the firewall logs, just shown in full. " + NOT_IN_LOGS, False), ("", False),
        (f"A2 ({len(model_rows)} rows): the model-review rows marked Unsure, with each firewall threat record and its signature,", True),
        ("   the biggest uploads, traffic per destination and the individual flows. Decide each again: Threat, Normal, Normal but", False),
        ("   unusual or Unsure. You are still not told which rows the model flagged.", False),
        (f"B2 ({len(blind_rows)} rows): the blind-check logs marked Unsure, with the same source's flows around each log.", True),
        (f"D2 ({len(devices)} rows): the campus devices behind the unidentified-app labels (group L5), one row per device, with", True),
        ("   its traffic summarised. Decide for each device's labels: Normal but unusual (not an attack), Keep as a threat, or Unsure.", False),
        ("   Your answers decide whether ATDR keeps raising alerts on unidentified-app traffic.", False),
        ("", False),
        ("Same rules as the pack: judge only from this file, do not look rows up in ATDR, and fill in 'Sign-off' honestly:", True),
        ("the method, the date you finished, and whether AI helped.", True), ("", False),
        *PRIVACY_LINES,
    ]
    _write_workbook(FOLLOWUP, guide=guide, signoff=SIGNOFF, sheets=[
        {"title": SHEET_A, "columns": MODEL_DETAIL_COLUMNS, "rows": model_rows, "choices": {"decision": DECISIONS, "confidence": CONFIDENCE}},
        {"title": SHEET_B, "columns": BLIND_DETAIL_COLUMNS, "rows": blind_rows, "choices": {"decision": DECISIONS, "confidence": CONFIDENCE}},
        {"title": SHEET_D, "columns": DEVICE_COLUMNS, "rows": devices, "choices": {"decision": DEVICE_CHOICES}},
    ])
    names.save(KEY)
    FOLLOWUP_KEY.write_text(json.dumps({"built_at": datetime.now(UTC).isoformat(), "model_rows": model_ids, "blind_rows": blind_ids,
                                        "devices": device_key}, indent=1), encoding="utf-8")
    print(f"wrote {FOLLOWUP}")
    print(f"  A2 {len(model_rows)} model rows, B2 {len(blind_rows)} blind-check rows, D2 {len(devices)} devices "
          f"({sum(len(ids) for ids in device_key.values())} labels)")
    print(f"private: {FOLLOWUP_KEY} (never send it)")


def read(pack: Path) -> None:
    from openpyxl import load_workbook

    key = json.loads(FOLLOWUP_KEY.read_text(encoding="utf-8"))
    pack_key = json.loads(PACK_KEY.read_text(encoding="utf-8"))
    workbook = load_workbook(pack, read_only=True, data_only=True)
    stamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")

    model_answers = {row["review_id"]: row for row in _sheet_rows(workbook, SHEET_A, MODEL_DETAIL_COLUMNS) if row.get("review_id")}
    model_path = PACK_DIR / "model_review_decisions.csv"
    shutil.copy(model_path, PACK_DIR / f"model_review_decisions.before_followup_{stamp}.csv")
    merged, model_changed = merge_followup(list(csv.DictReader(model_path.open(encoding="utf-8"))), model_answers, "review_id")
    _write_csv(model_path, merged, ["review_id", "decision", "confidence", "note"])

    blind_answers = {row["sample_id"]: row for row in _sheet_rows(workbook, SHEET_B, BLIND_DETAIL_COLUMNS) if row.get("sample_id")}
    blind_path = PACK_DIR / "blind_human_decisions.csv"
    shutil.copy(blind_path, PACK_DIR / f"blind_human_decisions.before_followup_{stamp}.csv")
    blind_merged, blind_changed = merge_followup(list(csv.DictReader(blind_path.open(encoding="utf-8"))), blind_answers, "sample_id")
    _write_csv(blind_path, blind_merged, ["sample_id", "decision", "confidence", "note"])
    ai_rows = list(csv.DictReader((BLIND / "decisions.csv").open(encoding="utf-8")))
    human = {row["sample_id"]: row for row in blind_merged}
    ai = {row["sample_id"]: row["decision"] for row in ai_rows if row["sample_id"] in pack_key["blind"]["rows"]}
    comparison = blind_comparison(ai, {sample: row["decision"] for sample, row in human.items()}, pack_key["blind"]["groups"])
    (PACK_DIR / "blind_comparison.json").write_text(json.dumps(comparison, indent=1), encoding="utf-8")
    checked = human_checked_decisions(ai_rows, human)
    _write_csv(PACK_DIR / "blind_decisions_human_checked.csv", checked, ["sample_id", "decision", "confidence", "note", "labeled_by"])
    human_score = score_blind_check(json.loads((BLIND / "key.json").read_text(encoding="utf-8")),
                                    {row["sample_id"]: row["decision"] for row in checked})
    (PACK_DIR / "blind_score_human_checked.json").write_text(json.dumps(human_score, indent=1, default=str), encoding="utf-8")

    devices = {row["device"]: row for row in _sheet_rows(workbook, SHEET_D, DEVICE_COLUMNS) if row.get("device")}
    group = {"key": "unidentified_app", "proposal": {"decision": "Normal but unusual", "attack_type": None},
             "log_ids": [log_id for ids in key["devices"].values() for log_id in ids]}
    with SessionLocal() as db:
        failed = pattern_condition_failures(db, {"D2": group})["D2"]
    decisions = []
    for device, log_ids in key["devices"].items():
        answer = devices.get(device) or {}
        normal = answer.get("decision") == DEVICE_CHOICES[0]
        decisions.append({"source_ip": "", "pattern": f"D2 unidentified-app device {device}",
                          "decision": "Normal but unusual" if normal else "Unsure", "note": answer.get("note") or "",
                          "log_ids": [log_id for log_id in log_ids if log_id not in failed]})
    signoff = {label: "" if value is None else str(value).strip()
               for label, value in workbook["Sign-off"].iter_rows(min_row=1, max_row=len(SIGNOFF), max_col=2, values_only=True)}
    problems = signoff_problems(signoff.get(SIGNOFF[2]), signoff.get(SIGNOFF[1]), signoff.get(SIGNOFF[4]))
    note = (f"team review follow-up, returned {datetime.now(UTC).date()}, reviewed by {signoff.get(SIGNOFF[0]) or 'unnamed'}, "
            f"method: {signoff.get(SIGNOFF[2]) or 'not stated'}")
    (PACK_DIR / "label_decisions_followup.json").write_text(json.dumps({"note": note, "decisions": decisions}, indent=1), encoding="utf-8")
    (PACK_DIR / "signoff_followup.json").write_text(json.dumps(signoff, indent=1, ensure_ascii=False), encoding="utf-8")

    print(f"A2: {len(model_changed)} of {len(key['model_rows'])} Unsure model rows now decided; the rest stay Unsure -> {model_path}")
    print(f"B2: {len(blind_changed)} of {len(key['blind_rows'])} Unsure blind logs now decided; blind check re-scored -> "
          f"{PACK_DIR / 'blind_score_human_checked.json'}")
    relabel = [item for item in decisions if item["decision"] != "Unsure"]
    print(f"D2: {len(relabel)} of {len(key['devices'])} devices called not an attack, covering {sum(len(item['log_ids']) for item in relabel)} labels"
          + (f" ({len(failed)} labels fail the condition and keep their label)" if failed else ""))
    print(f"Sign-off: {signoff.get(SIGNOFF[0]) or '(blank)'}; method {signoff.get(SIGNOFF[2]) or '(blank)'}; finished {signoff.get(SIGNOFF[1]) or '(blank)'}")
    print(f"  Note: {signoff.get(SIGNOFF[4]) or '(none)'}")
    print("  -> counts as a person's review" if not problems else "  -> does NOT count as a person's review: " + "; ".join(problems))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    commands = parser.add_subparsers(dest="command", required=True)
    commands.add_parser("build", help="write the follow-up and its private key")
    reader = commands.add_parser("read", help="merge a returned follow-up into the pack's outputs")
    reader.add_argument("--pack", type=Path, required=True)
    args = parser.parse_args()
    build() if args.command == "build" else read(args.pack)


if __name__ == "__main__":
    main()
