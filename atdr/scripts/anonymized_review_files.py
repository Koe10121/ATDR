"""Build review workbooks with MFU addresses replaced by consistent stand-ins.

Usage:
    python -m atdr.scripts.anonymized_review_files model-review [--round-dir .tmp/mfu_model/v2_fresh]
    python -m atdr.scripts.anonymized_review_files verify-blind-labels

model-review         the behaviour model's 22-row blind review (same row IDs as before)
verify-blind-labels  the team's hand check of the AI reviewer's blind-check labels: every
                     Threat, every false alarm and every Unsure, mixed with random other rows
                     so checkers cannot tell which rows ATDR flagged; plus the scoreboard's
                     threat-labeled BitTorrent logs that the file-sharing policy (v5.34.0)
                     no longer alerts on.

MFU addresses (private ones, and MFU public ones seen on an inside zone) become names such as
mfu-lan-3.12.7 that keep /16 and /24 grouping; internet addresses stay as they are. All files
share one key, .tmp/review_ip_key.json, which maps names back to real addresses and stays
private with the other files under .tmp/.
"""

from __future__ import annotations

import argparse
import csv
import json
import random
import re
import sqlite3
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ModuleNotFoundError as error:  # a team tool; ATDR itself does not need openpyxl
    raise SystemExit("This tool writes .xlsx files and needs openpyxl: pip install openpyxl") from error

from atdr.app.core.ip_pseudonyms import MfuPseudonyms, leftover_mfu_addresses
from atdr.app.detection.rules import INSIDE_ZONE_TOKENS, OUTSIDE_ZONE_TOKENS, P2P_SUBCATEGORY, P2P_TECHNOLOGY
from atdr.app.services.detection_scoreboard_service import THREAT_LABELS

ROOT = Path(__file__).resolve().parents[2]
TMP = ROOT / ".tmp"
KEY = TMP / "review_ip_key.json"
BLIND = TMP / "blind_check"
MODEL = TMP / "mfu_model"
LIVE_DB = ROOT / "atdr.db"
HOLDOUT_WINDOW = ("2026-05-20 13:45:00", "2026-05-20 13:57:00")
SEED = 20260928
DECOYS = 20
HARMLESS = {"Normal", "Normal but unusual"}

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
AI_FILL = PatternFill("solid", fgColor="4B3F72")
INPUT_HEADER_FILL = PatternFill("solid", fgColor="8A6D00")
INPUT_FILL = PatternFill("solid", fgColor="FFF2A8")
EXAMPLE_FILL = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="C8CED6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
IP_RANGE = re.compile(r"^\s*[0-9a-fA-F.:]+\s*-\s*[0-9a-fA-F.:]+\s*$")


# ------------------------------------------------------------------ MFU address space


def _inside(zone: str | None) -> bool:
    tokens = set(re.split(r"[^a-z0-9]+", (zone or "").lower()))
    return bool(tokens & INSIDE_ZONE_TOKENS) and not tokens & OUTSIDE_ZONE_TOKENS


def mfu_networks(database: Path, window: tuple[str, str] | None = None) -> list[str]:
    """Public /24 (IPv4) and /48 (IPv6) networks seen on an inside zone: MFU's own public ranges."""

    import ipaddress

    where = " WHERE generated_time >= ? AND generated_time < ?" if window else ""
    query = (f"SELECT DISTINCT src_ip, src_zone FROM normalized_logs{where} "
             f"UNION SELECT DISTINCT dst_ip, dst_zone FROM normalized_logs{where}")
    networks = set()
    with sqlite3.connect(f"file:{database.as_posix()}?mode=ro", uri=True) as connection:
        for address, zone in connection.execute(query, [*window, *window] if window else []):
            try:
                parsed = ipaddress.ip_address(str(address))
            except ValueError:
                continue
            if _inside(zone) and parsed.is_global:
                prefix = 24 if parsed.version == 4 else 48
                networks.add(str(ipaddress.ip_network(f"{parsed}/{prefix}", strict=False)))
    return sorted(networks)


def _anonymize_rows(rows: list[dict[str, Any]], names: MfuPseudonyms) -> list[dict[str, Any]]:
    cleaned = []
    for row in rows:
        out = {}
        for key, value in row.items():
            if isinstance(value, str) and key.endswith("country") and IP_RANGE.match(value):
                value = "private address range"
            out[key] = names.text(value)
        cleaned.append(out)
    leftovers = [value for row in cleaned for cell in row.values() if isinstance(cell, str)
                 for value in leftover_mfu_addresses(cell, names)]
    if leftovers:
        raise SystemExit(f"Refusing to write: {len(leftovers)} MFU addresses were not replaced.")
    return cleaned


# ------------------------------------------------------------------ workbook


def _write_workbook(path: Path, *, guide: list[tuple[str, bool]], sheets: list[dict[str, Any]], signoff: list[str]) -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "How to review"
    for index, (text, bold) in enumerate(guide, start=1):
        cell = first.cell(row=index, column=1, value=text)
        cell.font = Font(name=FONT, size=13 if index == 1 else 10, bold=bold)
    first.column_dimensions["A"].width = 120

    for spec in sheets:
        sheet = workbook.create_sheet(spec["title"])
        columns = spec["columns"]
        for column, (_key, title, width, kind) in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=column, value=title)
            cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=9)
            cell.fill = {"log": HEADER_FILL, "ai": AI_FILL, "input": INPUT_HEADER_FILL}[kind]
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            cell.border = BORDER
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.row_dimensions[1].height = 48
        body = [spec["example"], *spec["rows"]] if spec.get("example") else spec["rows"]
        for offset, values in enumerate(body):
            example = values is spec.get("example")
            for column, (key, _title, _width, kind) in enumerate(columns, start=1):
                cell = sheet.cell(row=2 + offset, column=column, value=values.get(key, ""))
                cell.font = Font(name=FONT, size=9, italic=example)
                cell.border = BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if kind == "input":
                    cell.fill = EXAMPLE_FILL if example else INPUT_FILL
                elif example:
                    cell.fill = EXAMPLE_FILL
        first_data = 3 if spec.get("example") else 2
        last = 1 + len(body)
        for key, options in spec["choices"].items():
            if any("," in option for option in options):
                raise ValueError(f"Excel splits dropdown choices on commas: {options}")
            letter = get_column_letter(next(i for i, column in enumerate(columns, start=1) if column[0] == key))
            rule = DataValidation(type="list", formula1=f'"{",".join(options)}"', allow_blank=True, showErrorMessage=True)
            sheet.add_data_validation(rule)
            rule.add(f"{letter}{first_data}:{letter}{last}")
        sheet.freeze_panes = "B2"

    sign = workbook.create_sheet("Sign-off")
    for index, label in enumerate(signoff, start=1):
        sign.cell(row=index, column=1, value=label).font = Font(name=FONT, bold=True, size=10)
        cell = sign.cell(row=index, column=2)
        cell.fill = INPUT_FILL
        cell.border = BORDER
    sign.column_dimensions["A"].width = 64
    sign.column_dimensions["B"].width = 60
    for row, options in SIGNOFF_CHOICES.items():
        if row <= len(signoff):
            rule = DataValidation(type="list", formula1=f'"{",".join(options)}"', allow_blank=True, showErrorMessage=True)
            sign.add_data_validation(rule)
            rule.add(f"B{row}")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


PRIVACY_LINES = [
    ("Addresses: MFU addresses are replaced by stand-ins such as mfu-lan-3.12.7 (private block 3, subnet 12, host 7) or", True),
    ("mfu-pub-1.2.5 (an MFU public address). The same address always has the same stand-in in every file, and hosts that share", False),
    ("a subnet share its prefix. Internet addresses are real, so you can still weigh who the outside party is.", False),
]
SIGNOFF = ["Reviewed by (names)", "Date finished", "How was it reviewed? (pick one)",
           "Did anyone look up these rows in ATDR, or learn which rows ATDR flagged?", "Anything we should know"]
# Only a review done or checked by a person can switch an attack type on (ML_QUALITY_BAR.md), so the
# sign-off must say which it was. No commas: Excel splits dropdown choices on them.
SIGNOFF_CHOICES = {3: ["By a person without AI", "AI-assisted and a person checked every row", "AI only (not checked by a person)"],
                   4: ["No", "Yes (explain below)"]}


# ------------------------------------------------------------------ model review


MODEL_COLUMNS = [
    ("review_id", "Row", 8, "log"), ("source", "Source", 17, "log"), ("from", "First seen", 10, "log"),
    ("to", "Last seen", 10, "log"), ("direction", "Direction (zone -> zone, connections)", 30, "log"),
    ("connections", "Connections in 5 min", 11, "log"), ("different_destinations", "Different destination IPs", 12, "log"),
    ("different_destination_ports", "Different destination ports", 12, "log"),
    ("top_destinations", "Busiest destinations (IP:port, connections)", 44, "log"), ("actions", "Firewall actions", 16, "log"),
    ("bytes_sent", "Bytes sent (total)", 14, "log"), ("bytes_received", "Bytes received (total)", 14, "log"),
    ("largest_upload", "Largest single upload (bytes)", 14, "log"), ("top_apps", "Apps (connections)", 32, "log"),
    ("steadiest_repeated_destination", "Most regular repeated destination", 44, "log"),
    ("firewall_threat_logs", "Firewall threat logs", 10, "log"),
    ("decision", "YOUR DECISION", 20, "input"), ("confidence", "Confidence", 12, "input"), ("note", "Note (why)", 40, "input"),
]
MODEL_EXAMPLE = {
    "review_id": "EXAMPLE", "source": "mfu-lan-1.4.9", "from": "13:50:03", "to": "13:54:58",
    "direction": "WLAN-Inside -> SG-Outside (40)", "connections": 40, "different_destinations": 2,
    "different_destination_ports": 2, "top_destinations": "185.x.x.x:8443 (30), 142.x.x.x:443 (10)",
    "actions": "allow 40", "bytes_sent": 36000, "bytes_received": 51000, "largest_upload": 1400,
    "top_apps": "unknown-tcp (30), ssl (10)",
    "steadiest_repeated_destination": "185.x.x.x: 30 connections, every 10s on average (variation 0.04)",
    "firewall_threat_logs": 0, "decision": "Threat", "confidence": "Medium",
    "note": "Laptop calls one unknown server every 10 s like clockwork on an odd port: looks like malware checking in. (Example only.)",
}


def _period(round_dir: Path) -> str:
    key = json.loads((round_dir / "review_key.json").read_text(encoding="utf-8"))
    windows = sorted({entry["window"][11:16] for entry in key["reviews"]})
    return "20 May, windows starting " + ", ".join(windows)


def build_model_review(names: MfuPseudonyms, round_dir: Path = MODEL) -> Path:
    rows = _anonymize_rows(json.loads((round_dir / "review_sample.json").read_text(encoding="utf-8")), names)
    out = round_dir / ("ATDR_model_review_anonymized.xlsx" if round_dir == MODEL else f"ATDR_model_review_{round_dir.name}_anonymized.xlsx")
    guide = [
        (f"ATDR model review: {len(rows)} short rows (anonymized)", True), ("", False),
        (f"Each row is what one device did during five minutes ({_period(round_dir)}), a part of the MFU file that ATDR", False),
        ("never trained on. Some rows are things a new detection model flagged; others are random. You are not told which.", False),
        ("Please judge every row the same way.", False), ("", False),
        ("Judge only from this sheet (plus general knowledge or outside lookups of internet addresses). Do not look rows up", True),
        ("in ATDR, and do not ask Claude which rows were flagged.", True), ("", False),
        ("YOUR DECISION: Threat, Normal, Normal but unusual, or Unsure. Peer-to-peer file sharing (BitTorrent and the like) is", False),
        ("policy activity for ATDR: call it Normal but unusual unless there is other evidence of an attack.", False),
        ("'Most regular repeated destination' shows a device contacting the same server at steady intervals (malware often", False),
        ("checks in like that; so do some normal apps). 'Largest single upload' shows big data leaving the device.", False), ("", False),
        *PRIVACY_LINES, ("", False),
        ("The grey EXAMPLE row shows the format and is not scored. Fill in 'Sign-off' when done and send the file back.", False),
        ("On 'Sign-off', say how the review was done. An attack type can only be switched on from a review done or checked", True),
        ("by a person, so if AI helped, a person must check every row before signing.", True),
    ]
    _write_workbook(out, guide=guide, signoff=SIGNOFF, sheets=[{
        "title": "Review", "columns": MODEL_COLUMNS, "rows": rows, "example": MODEL_EXAMPLE,
        "choices": {"decision": ["Threat", "Normal", "Normal but unusual", "Unsure"], "confidence": ["High", "Medium", "Low"]},
    }])
    return out


# ------------------------------------------------------------------ blind label verification


BLIND_COLUMNS = [
    ("sample_id", "Sample", 9, "log"), ("time", "Time (20 May)", 11, "log"), ("log_type", "Log type", 13, "log"),
    ("action", "Firewall action", 13, "log"), ("source", "Source IP:port", 22, "log"), ("source_zone", "Source zone", 12, "log"),
    ("source_country", "Source country", 16, "log"), ("destination", "Destination IP:port", 24, "log"),
    ("destination_zone", "Destination zone", 13, "log"), ("destination_country", "Destination country", 16, "log"),
    ("protocol", "Protocol", 9, "log"), ("app", "App", 16, "log"), ("app_category", "App category", 22, "log"),
    ("app_risk", "App risk (1-5)", 9, "log"), ("bytes_sent", "Bytes sent", 11, "log"), ("bytes_received", "Bytes received", 12, "log"),
    ("packets", "Packets", 9, "log"), ("elapsed_seconds", "Session length (s)", 10, "log"),
    ("session_end_reason", "Session end reason", 16, "log"), ("firewall_rule", "Firewall rule", 18, "log"),
    ("firewall_threat", "Firewall threat (name / severity / direction)", 30, "log"),
    ("source_logs_5min", "Same source: logs within 2.5 min", 12, "log"),
    ("source_distinct_destinations", "Same source: different destination IPs", 13, "log"),
    ("source_distinct_ports", "Same source: different destination ports", 13, "log"),
    ("source_denied_or_reset", "Same source: denied or reset", 12, "log"), ("source_total_bytes", "Same source: total bytes", 13, "log"),
    ("source_top_apps", "Same source: top apps (logs)", 32, "log"),
    ("sources_to_same_destination_port", "Different sources hitting the same destination:port", 14, "log"),
    ("ai_decision", "AI reviewer's decision", 16, "ai"), ("ai_confidence", "AI confidence", 10, "ai"), ("ai_note", "AI reviewer's note", 44, "ai"),
    ("verdict", "AGREE OR CHANGE?", 14, "input"), ("your_decision", "Your decision (if changed)", 18, "input"), ("note", "Note (why)", 40, "input"),
]
TORRENT_COLUMNS = [
    ("row", "Row", 7, "log"), ("time", "Time (20 May)", 11, "log"), ("source", "Source", 17, "log"),
    ("destination", "Destination IP:port", 24, "log"), ("direction", "Zones", 22, "log"), ("app", "App", 12, "log"),
    ("action", "Firewall action", 11, "log"), ("bytes", "Bytes sent / received", 16, "log"),
    ("session_end_reason", "Session end reason", 14, "log"),
    ("source_logs", "Same source: logs in the data", 11, "log"), ("source_p2p_share", "Same source: share that is file sharing", 12, "log"),
    ("source_destinations", "Same source: different destination IPs", 12, "log"),
    ("source_ports", "Same source: different destination ports", 12, "log"),
    ("source_threat_logs", "Same source: firewall threat logs", 11, "log"),
    ("team_label", "Team's label", 20, "ai"), ("label_source", "How it was labeled", 14, "ai"), ("team_note", "Label note", 44, "ai"),
    ("verdict", "STILL A THREAT?", 18, "input"), ("note", "Note (why)", 40, "input"),
]


def _torrent_rows() -> list[dict[str, Any]]:
    """Threat-labeled file-sharing logs whose attack type is not a policy violation."""

    query = """
        WITH latest AS (
            SELECT l.* FROM ml_labels l
            WHERE l.reviewed = 1 AND l.id = (SELECT MAX(id) FROM ml_labels WHERE log_id = l.log_id AND reviewed = 1)
        )
        SELECT n.generated_time, n.src_ip, n.dst_ip, n.dst_port, n.src_zone, n.dst_zone, n.app, n.action,
               n.bytes_sent, n.bytes_received, n.session_end_reason, l.label, l.attack_type, l.label_source, l.review_note
        FROM latest l JOIN normalized_logs n ON n.id = l.log_id
        WHERE lower(n.app_technology) = ? AND lower(n.app_subcategory) = ? AND l.attack_type != 'policy_violation'
        ORDER BY n.generated_time
    """
    rows = []
    with sqlite3.connect(f"file:{LIVE_DB.as_posix()}?mode=ro", uri=True) as connection:
        found = [row for row in connection.execute(query, (P2P_TECHNOLOGY, P2P_SUBCATEGORY)) if row[11] in THREAT_LABELS]
        context = {}
        for src in {row[1] for row in found}:
            context[src] = connection.execute(
                "SELECT COUNT(*), AVG(lower(app_technology) = ? AND lower(app_subcategory) = ?), COUNT(DISTINCT dst_ip), "
                "COUNT(DISTINCT dst_port), SUM(upper(log_type) = 'THREAT') FROM normalized_logs WHERE src_ip = ?",
                (P2P_TECHNOLOGY, P2P_SUBCATEGORY, src),
            ).fetchone()
    for index, row in enumerate(found, start=1):
        (when, src, dst, port, src_zone, dst_zone, app, action, sent, received, end, label, attack, source, note) = row
        logs, share, destinations, ports, threats = context[src]
        rows.append({
            "row": f"T{index:02d}", "time": str(when)[11:19], "source": src, "destination": f"{dst}:{port}",
            "direction": f"{src_zone} -> {dst_zone}", "app": app, "action": action, "bytes": f"{sent or 0:,.0f} / {received or 0:,.0f}",
            "session_end_reason": end, "source_logs": logs, "source_p2p_share": f"{(share or 0):.0%}",
            "source_destinations": destinations, "source_ports": ports, "source_threat_logs": threats or 0,
            "team_label": f"{label} / {attack}", "label_source": source, "team_note": note or "",
        })
    return rows


def build_blind_verification(names: MfuPseudonyms) -> Path:
    samples = {row["sample_id"]: row for row in json.loads((BLIND / "sample.json").read_text(encoding="utf-8"))}
    key = json.loads((BLIND / "key.json").read_text(encoding="utf-8"))
    decisions = {row["sample_id"]: row for row in csv.DictReader((BLIND / "decisions.csv").open(encoding="utf-8"))}
    model_flags = set(json.loads((MODEL / "evaluation.json").read_text(encoding="utf-8"))["model_flagged_log_ids"])
    flagged = {entry["sample_id"] for entry in key["samples"] if entry["alerted"] or entry["log_id"] in model_flags}

    groups: dict[str, list[str]] = {"threat": [], "false_alarm": [], "unsure": [], "other": []}
    for sample_id, decision in decisions.items():
        verdict = decision["decision"]
        group = ("threat" if verdict == "Threat" else "unsure" if verdict == "Unsure"
                 else "false_alarm" if sample_id in flagged and verdict in HARMLESS else "other")
        groups[group].append(sample_id)
    rng = random.Random(SEED)
    decoys = rng.sample(sorted(groups["other"]), min(DECOYS, len(groups["other"])))
    chosen = [*groups["threat"], *groups["false_alarm"], *groups["unsure"], *decoys]
    rng.shuffle(chosen)

    raw_rows = []
    for sample_id in chosen:
        decision = decisions[sample_id]
        raw_rows.append({**samples[sample_id], "ai_decision": decision["decision"], "ai_confidence": decision["confidence"],
                         "ai_note": decision["note"]})
    rows = _anonymize_rows(raw_rows, names)
    torrents = _anonymize_rows(_torrent_rows(), names)
    selection = {"seed": SEED, "rows": len(chosen), "groups": {name: sorted(ids) for name, ids in groups.items() if name != "other"},
                 "decoys": sorted(decoys)}
    (BLIND / "verification_selection.json").write_text(json.dumps(selection, indent=1), encoding="utf-8")

    out = BLIND / "ATDR_blind_label_verification.xlsx"
    guide = [
        ("ATDR blind check: human verification of the AI reviewer's labels", True), ("", False),
        (f"Sheet 'Blind labels' ({len(rows)} rows): logs from the blind check. An AI reviewer (ChatGPT, acting as a senior SOC analyst)", False),
        ("labeled them without seeing anything from ATDR. Please check each of its decisions: pick Agree, or pick Change and give", False),
        ("your decision. The rows are a mix; you are not told why each one is here, so judge every row the same way.", False),
        ("Decisions: Threat, Normal, Normal but unusual, or Unsure, the same meanings as the blind-check sheet.", False),
        ("Peer-to-peer file sharing (BitTorrent and the like) is policy activity for ATDR: Normal but unusual unless there is", False),
        ("other evidence of an attack. A single unanswered probe from the internet is background noise, not a Threat on its own.", False),
        ("", False),
        (f"Sheet 'BitTorrent labels' ({len(torrents)} rows): logs our own team labeled as threats (mostly port scans) that are", False),
        ("BitTorrent file sharing. ATDR no longer alerts on file sharing without malicious evidence. For each: is it still a", False),
        ("threat, or is it policy activity (file sharing's normal fan-out)?", False),
        ("", False),
        ("Do not look rows up in ATDR and do not ask Claude which rows ATDR flagged.", True), ("", False),
        *PRIVACY_LINES,
    ]
    _write_workbook(out, guide=guide, signoff=SIGNOFF, sheets=[
        {"title": "Blind labels", "columns": BLIND_COLUMNS, "rows": rows,
         "choices": {"verdict": ["Agree", "Change"], "your_decision": ["Threat", "Normal", "Normal but unusual", "Unsure"]}},
        {"title": "BitTorrent labels", "columns": TORRENT_COLUMNS, "rows": torrents,
         "choices": {"verdict": ["Still a threat", "Policy activity (not a threat)", "Unsure"]}},
    ])
    return out


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("what", choices=["model-review", "verify-blind-labels"])
    parser.add_argument("--round-dir", type=Path, default=MODEL, help="model review round (default: the v1 review)")
    args = parser.parse_args()
    networks = sorted(set(mfu_networks(BLIND / "holdout.db", HOLDOUT_WINDOW)) | set(mfu_networks(LIVE_DB)))
    names = MfuPseudonyms.load(KEY, networks)
    out = build_model_review(names, args.round_dir) if args.what == "model-review" else build_blind_verification(names)
    names.save(KEY)
    print(f"wrote {out}")
    print(f"MFU addresses named so far: {len(names.mapping)} (key: {KEY}, private)")


if __name__ == "__main__":
    main()
