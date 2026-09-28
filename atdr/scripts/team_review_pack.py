"""One workbook with everything ATDR still needs a person to judge, and a reader for the returned copy.

Usage:
    python -m atdr.scripts.team_review_pack build
    python -m atdr.scripts.team_review_pack read --pack "returned pack.xlsx"

build  writes .tmp/team_review/ATDR_team_review_pack.xlsx (MFU addresses anonymized, same stand-ins
       as every earlier review file) and the private pack_key.json that maps its rows back.
read   turns a returned pack into input for the existing tools, all under .tmp/team_review/:
       model_review_decisions.csv  -> evaluate_behavior_model --review-decisions (part A)
       blind_human_decisions.csv, blind_comparison.json, blind_decisions_human_checked.csv and
       blind_score_human_checked.json: the blind check re-scored with the team's labels (part B;
       the official .tmp/blind_check/score.json is never touched)
       label_decisions.json        -> apply_label_review (parts C and D)
       signoff.json
       It changes nothing in ATDR; the printed commands apply the results.

Parts: A. the MFU behaviour model's extra finds, B. the blind-check logs judged without the AI
reviewer's decision, C. BitTorrent logs the team labeled as threats, D. groups of team labels that
contradict the firewall's evidence or the team's policies. The logic is in
atdr/app/services/team_review_service.py; this tool needs openpyxl, like anonymized_review_files.
"""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from atdr.app.core.ip_pseudonyms import MfuPseudonyms
from atdr.app.db.database import SessionLocal
from atdr.app.services.blind_check_service import score_blind_check
from atdr.app.services.team_review_service import (
    PATTERN_CHOICES,
    TORRENT_POLICY,
    blind_comparison,
    example_text,
    human_checked_decisions,
    label_patterns,
    pattern_condition_failures,
    pattern_label_decisions,
    signoff_problems,
    torrent_label_decisions,
)
from atdr.scripts.anonymized_review_files import (
    BLIND,
    BLIND_COLUMNS,
    HOLDOUT_WINDOW,
    KEY,
    LIVE_DB,
    MODEL,
    MODEL_COLUMNS,
    MODEL_EXAMPLE,
    PRIVACY_LINES,
    SIGNOFF,
    TMP,
    TORRENT_COLUMNS,
    _anonymize_rows,
    _torrent_rows,
    _write_workbook,
    blind_selection,
    mfu_networks,
)

PACK_DIR = TMP / "team_review"
PACK = PACK_DIR / "ATDR_team_review_pack.xlsx"
PACK_KEY = PACK_DIR / "pack_key.json"
MODEL_ROUND = MODEL / "v2_fresh"
DECISIONS = ["Threat", "Normal", "Normal but unusual", "Unsure"]
CONFIDENCE = ["High", "Medium", "Low"]

SHEET_A, SHEET_B, SHEET_C, SHEET_D = "A. Model finds", "B. Blind-check logs", "C. BitTorrent labels", "D. Label patterns"
# Part B hides the AI reviewer's decision: the person decides from scratch and the two are compared.
BLIND_TEAM_COLUMNS = [column for column in BLIND_COLUMNS if column[3] == "log"] + [
    ("decision", "YOUR DECISION", 18, "input"), ("confidence", "Confidence", 12, "input"), ("note", "Note (why)", 40, "input"),
]
PATTERN_COLUMNS = [
    ("id", "Group", 7, "log"), ("title", "What the labels are", 26, "log"), ("what", "Why we ask", 58, "log"),
    ("logs", "Logs", 7, "log"), ("sources", "Sources", 8, "log"), ("current_labels", "Current labels (logs)", 32, "log"),
    ("labeled_by", "How they were labeled", 30, "log"), ("examples", "First examples", 70, "log"),
    ("proposal_text", "Proposed label", 24, "log"), ("choice", "YOUR DECISION", 22, "input"), ("note", "Note (why)", 40, "input"),
]


def _names() -> MfuPseudonyms:
    networks = sorted(set(mfu_networks(BLIND / "holdout.db", HOLDOUT_WINDOW)) | set(mfu_networks(LIVE_DB)))
    return MfuPseudonyms.load(KEY, networks)


def build() -> None:
    names = _names()
    model_rows = _anonymize_rows(json.loads((MODEL_ROUND / "review_sample.json").read_text(encoding="utf-8")), names)

    selection = blind_selection()
    blind_rows = _anonymize_rows([selection["samples"][sample] for sample in selection["chosen"]], names)

    torrents = _torrent_rows()
    torrent_key = {row["row"]: row["log_id"] for row in torrents}
    torrent_rows = _anonymize_rows([{key: value for key, value in row.items() if key != "log_id"} for row in torrents], names)

    with SessionLocal() as db:
        patterns = label_patterns(db)
    pattern_rows = _anonymize_rows([
        {"id": pattern["id"], "title": pattern["title"], "what": pattern["what"], "logs": pattern["logs"],
         "sources": pattern["sources"], "current_labels": pattern["current_labels"], "labeled_by": pattern["labeled_by"],
         "examples": "\n".join(example_text(example) for example in pattern["examples"]), "proposal_text": pattern["proposal_text"]}
        for pattern in patterns
    ], names)

    guide = [
        ("ATDR team review pack: the checks that need a person", True), ("", False),
        ("ATDR's accuracy figures and its MFU-trained model still rest on labels made by an AI reviewer or by ATDR's own rules.", False),
        ("These four parts are the human check. Together about two hours for one person; the parts can be split between people.", False),
        ("", False),
        (f"A. Model finds ({len(model_rows)} rows, about 30 min). What one device did in five minutes of traffic the model never", True),
        ("   trained on. Some rows are things the MFU model flagged, others are random; you are not told which. Decide each row:", False),
        ("   Threat, Normal, Normal but unusual, or Unsure. This decides whether the model may raise alerts for port scans.", False),
        (f"B. Blind-check logs ({len(blind_rows)} rows, about an hour). Single firewall logs with context about their source. Decide", True),
        ("   each yourself with the same four choices. An AI reviewer judged them before; its answer is hidden on purpose, and", False),
        ("   ATDR compares the two afterwards. Your decisions turn the blind check into human-checked labels.", False),
        (f"C. BitTorrent labels ({len(torrent_rows)} rows, about 10 min). Logs our team labeled as threats that are file sharing.", True),
        ("   ATDR now treats file sharing as policy activity unless there is other evidence. Still a threat, or policy activity?", False),
        (f"D. Label patterns ({len(pattern_rows)} rows, about 15 min). Each row is a group of our own labels that contradicts the", True),
        ("   firewall's own threat record or a policy the team agreed. Read the examples, then pick 'Relabel as proposed',", False),
        ("   'Keep the current label' or 'Unsure'. Nothing changes until the team lead applies the returned file.", False),
        ("   One group (campus devices using unidentified apps) also decides whether ATDR keeps alerting on that traffic.", False),
        ("", False),
        ("Peer-to-peer file sharing is policy activity: Normal but unusual unless there is other evidence of an attack.", False),
        ("A single unanswered probe from the internet is background noise, not a Threat on its own.", False),
        ("", False),
        ("Judge only from this file (outside lookups of internet addresses are fine). Do not look rows up in ATDR, and do not", True),
        ("ask Claude, ChatGPT or ATDR's assistant which rows were flagged. If AI helps, a person must check every row and", True),
        ("say so on 'Sign-off': only a review done or checked by a person counts.", True),
        ("", False),
        *PRIVACY_LINES, ("", False),
        ("The grey EXAMPLE row in part A shows the format and is not scored. Fill in 'Sign-off' when done and send the file back.", False),
        ("'Sign-off' counts only with a person's method, the date you finished, and no note saying it is an AI draft.", True),
    ]
    _write_workbook(PACK, guide=guide, signoff=SIGNOFF, sheets=[
        {"title": SHEET_A, "columns": MODEL_COLUMNS, "rows": model_rows, "example": MODEL_EXAMPLE,
         "choices": {"decision": DECISIONS, "confidence": CONFIDENCE}},
        {"title": SHEET_B, "columns": BLIND_TEAM_COLUMNS, "rows": blind_rows,
         "choices": {"decision": DECISIONS, "confidence": CONFIDENCE}},
        {"title": SHEET_C, "columns": TORRENT_COLUMNS, "rows": torrent_rows,
         "choices": {"verdict": ["Still a threat", TORRENT_POLICY, "Unsure"]}},
        {"title": SHEET_D, "columns": PATTERN_COLUMNS, "rows": pattern_rows, "choices": {"choice": list(PATTERN_CHOICES)}},
    ])
    names.save(KEY)
    groups = {**{name: sorted(ids) for name, ids in selection["groups"].items() if name != "other"},
              "decoy": sorted(selection["decoys"])}
    key = {
        "built_at": datetime.now(UTC).isoformat(),
        "model_round": str(MODEL_ROUND.relative_to(TMP.parent)),
        "blind": {"rows": selection["chosen"], "groups": groups},
        "torrent": torrent_key,
        "patterns": {pattern["id"]: {"key": pattern["key"], "title": pattern["title"], "proposal": pattern["proposal"],
                                     "log_ids": pattern["log_ids"]} for pattern in patterns},
    }
    PACK_KEY.write_text(json.dumps(key, indent=1), encoding="utf-8")
    print(f"wrote {PACK}")
    print(f"  A {len(model_rows)} model rows, B {len(blind_rows)} blind-check rows, C {len(torrent_rows)} BitTorrent rows, "
          f"D {len(pattern_rows)} label groups ({sum(p['logs'] for p in patterns)} labels)")
    print(f"private: {PACK_KEY} and {KEY} (never send these)")


# ------------------------------------------------------------------ read


def _sheet_rows(workbook: Any, title: str, columns: list[tuple]) -> list[dict[str, str]]:
    sheet = workbook[title]
    by_title = {column[1]: column[0] for column in columns}
    keys = [by_title.get(str(cell.value or "")) for cell in sheet[1]]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = {key: "" if value is None else str(value).strip() for key, value in zip(keys, values) if key}
        if any(row.values()):
            rows.append(row)
    return rows


def _write_csv(path: Path, rows: list[dict[str, Any]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def read(pack: Path, *, method: str | None = None, method_reason: str | None = None) -> None:
    from openpyxl import load_workbook

    key = json.loads(PACK_KEY.read_text(encoding="utf-8"))
    workbook = load_workbook(pack, read_only=True, data_only=True)
    PACK_DIR.mkdir(parents=True, exist_ok=True)

    model = [row for row in _sheet_rows(workbook, SHEET_A, MODEL_COLUMNS) if row.get("review_id") not in ("", "EXAMPLE")]
    model_done = [row for row in model if row.get("decision")]
    _write_csv(PACK_DIR / "model_review_decisions.csv", model_done, ["review_id", "decision", "confidence", "note"])

    blind = {row["sample_id"]: row for row in _sheet_rows(workbook, SHEET_B, BLIND_TEAM_COLUMNS) if row.get("sample_id")}
    ai_rows = list(csv.DictReader((BLIND / "decisions.csv").open(encoding="utf-8")))
    ai = {row["sample_id"]: row["decision"] for row in ai_rows if row["sample_id"] in key["blind"]["rows"]}
    comparison = blind_comparison(ai, {sample: row.get("decision", "") for sample, row in blind.items()}, key["blind"]["groups"])
    (PACK_DIR / "blind_comparison.json").write_text(json.dumps(comparison, indent=1), encoding="utf-8")
    _write_csv(PACK_DIR / "blind_human_decisions.csv", [row for row in blind.values() if row.get("decision")],
               ["sample_id", "decision", "confidence", "note"])
    checked = human_checked_decisions(ai_rows, blind)
    _write_csv(PACK_DIR / "blind_decisions_human_checked.csv", checked, ["sample_id", "decision", "confidence", "note", "labeled_by"])
    # Same method as the official score (the rules' verdicts recorded in the key), labels checked by a person.
    blind_key = json.loads((BLIND / "key.json").read_text(encoding="utf-8"))
    human_score = score_blind_check(blind_key, {row["sample_id"]: row["decision"] for row in checked})
    (PACK_DIR / "blind_score_human_checked.json").write_text(json.dumps(human_score, indent=1, default=str), encoding="utf-8")

    torrents = {row["row"]: row for row in _sheet_rows(workbook, SHEET_C, TORRENT_COLUMNS) if row.get("row")}
    patterns = {row["id"]: row for row in _sheet_rows(workbook, SHEET_D, PATTERN_COLUMNS) if row.get("id")}
    # Relabel only the labels that meet each group's condition; the others keep their label and are listed.
    torrent_group = {"key": "file_sharing", "proposal": {"decision": "Normal but unusual", "attack_type": None},
                     "log_ids": list(key["torrent"].values())}
    with SessionLocal() as db:
        failures = pattern_condition_failures(db, key["patterns"])
        torrent_failures = pattern_condition_failures(db, {"C": torrent_group})["C"]
    for row, log_id in key["torrent"].items():
        if log_id in torrent_failures and (torrents.get(row) or {}).get("verdict") == TORRENT_POLICY:
            print(f"   {row}: keeps its label ({torrent_failures[log_id]})")
            torrents[row] = {**torrents[row], "verdict": "Unsure"}
    checked_patterns = {pattern_id: {**pattern, "log_ids": [log_id for log_id in pattern["log_ids"] if log_id not in failures[pattern_id]]}
                        for pattern_id, pattern in key["patterns"].items()}
    (PACK_DIR / "label_conditions.json").write_text(json.dumps(
        {pattern_id: {"labels": len(key["patterns"][pattern_id]["log_ids"]), "meeting_condition": len(checked_patterns[pattern_id]["log_ids"]),
                      "kept_as_labeled": {str(log_id): reason for log_id, reason in failed.items()}}
         for pattern_id, failed in failures.items()}, indent=1), encoding="utf-8")
    decisions = pattern_label_decisions(checked_patterns, patterns) + torrent_label_decisions(key["torrent"], torrents)
    signoff_sheet = workbook["Sign-off"]
    signoff = {label: "" if value is None else str(value).strip()
               for label, value in signoff_sheet.iter_rows(min_row=1, max_row=len(SIGNOFF), max_col=2, values_only=True)}
    if method:
        signoff["Method as picked on the sign-off"] = signoff.get(SIGNOFF[2], "")
        signoff[SIGNOFF[2]] = method
        signoff["Why the method was corrected"] = method_reason or ""
    method = signoff.get(SIGNOFF[2], "")
    note = f"team review pack, returned {datetime.now(UTC).date()}, reviewed by {signoff.get(SIGNOFF[0]) or 'unnamed'}, method: {method or 'not stated'}"
    (PACK_DIR / "label_decisions.json").write_text(json.dumps({"note": note, "decisions": decisions}, indent=1), encoding="utf-8")
    (PACK_DIR / "signoff.json").write_text(json.dumps(signoff, indent=1, ensure_ascii=False), encoding="utf-8")

    problems = signoff_problems(method, signoff.get(SIGNOFF[1]), signoff.get(SIGNOFF[4]))
    person = not problems
    print(f"A  model finds: {len(model_done)} of {len(model)} decided -> {PACK_DIR / 'model_review_decisions.csv'}")
    overall = comparison["overall"]
    print(f"B  blind-check logs: {comparison['judged']} of {len(key['blind']['rows'])} decided; same decision as the AI reviewer "
          f"{overall['same_decision']}/{overall['rows']}, same threat-or-not {overall['same_threat_or_not']}/{overall['threat_or_not_rows']}")
    estimate = human_score["estimate"]
    print("   blind check re-scored with these labels (rules v5.32.0, as the official result): "
          + ", ".join(f"{name.replace('_', ' ')} {value * 100:.1f}%" for name, value in estimate.items() if value is not None)
          + f" -> {PACK_DIR / 'blind_score_human_checked.json'}")
    for pattern_id, failed in failures.items():
        answer = (patterns.get(pattern_id) or {}).get("choice")
        if failed and answer == PATTERN_CHOICES[0]:
            reasons = Counter(failed.values()).most_common(2)
            print(f"   {pattern_id}: {len(failed)} of {len(key['patterns'][pattern_id]['log_ids'])} labels fail the group's condition and keep "
                  f"their label ({'; '.join(f'{reason} ({count})' for reason, count in reasons)})")
    relabel = [item for item in decisions if item["decision"] != "Unsure"]
    print(f"C+D label changes proposed by the team: {len(relabel)} decisions covering {sum(len(item['log_ids']) for item in relabel)} labels")
    print(f"Sign-off: reviewed by {signoff.get(SIGNOFF[0]) or '(blank)'}; method: {method or '(blank)'}; "
          f"finished: {signoff.get(SIGNOFF[1]) or '(blank)'}")
    print(f"  Note: {signoff.get(SIGNOFF[4]) or '(none)'}")
    if person:
        print("  -> counts as a person's review")
    else:
        print("  -> does NOT count as a person's review: " + "; ".join(problems))
        print("  Do not apply these results: the model review cannot switch a type on, and the labels stay")
        print("  'independent blind reference labeling'. The commands below are for a person-checked return only.")
    print("\nNext (each prints its effect before changing anything):")
    print(f"  python -m atdr.scripts.evaluate_behavior_model --out-dir {key['model_round']} "
          f"--review-decisions {PACK_DIR.relative_to(TMP.parent) / 'model_review_decisions.csv'} --min-reviewed 5")
    print(f"  python -m atdr.scripts.evaluate_behavior_model --out-dir {key['model_round']} --record-bar "
          f"--blind-dir .tmp/mfu_model/v2_second_look --min-reviewed 5 --review-method \"{method}\"")
    print(f"  python -m atdr.scripts.apply_label_review {PACK_DIR.relative_to(TMP.parent) / 'label_decisions.json'} "
          f"--reviewer \"{signoff.get(SIGNOFF[0]) or 'team'}\"   (add --apply to write)")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    commands = parser.add_subparsers(dest="command", required=True)
    commands.add_parser("build", help="write the pack and its private key")
    reader = commands.add_parser("read", help="turn a returned pack into input for the existing tools")
    reader.add_argument("--pack", type=Path, required=True)
    reader.add_argument("--method", help="the review method as confirmed by the team lead, when the sign-off picked another")
    reader.add_argument("--method-reason", help="why the sign-off's method was corrected (recorded with it)")
    args = parser.parse_args()
    if args.command == "build":
        build()
    else:
        if args.method and not args.method_reason:
            parser.error("--method needs --method-reason")
        read(args.pack, method=args.method, method_reason=args.method_reason)


if __name__ == "__main__":
    main()
